#!/usr/bin/env python3
"""Create the editable, macro-free ANA CRM training presentation."""
from pathlib import Path
import json
import hashlib
import math
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'artifacts'
CAP = ROOT / 'docs/capacitacion/capturas'
LOGO = ROOT / 'src/assets/ana-logo.png'
OUT.mkdir(exist_ok=True)
W, H = 13.333333, 7.5
C = dict(blue='2957A4', navy='17345E', ink='1D3555', muted='617793',
         light='EEF4FC', line='D8E3F3', white='FFFFFF', aqua='E6F2FA',
         green='22705C', green_bg='EDF7F3', yellow='FFF3CD', ochre='8A610E')
APP = 'https://crm-whatsapp-ana.pages.dev'
REPO = 'https://github.com/GabrielCuba23/CRM_ANA'
ZIP = REPO + '/raw/refs/heads/main/artifacts/crm-ana-extension.zip'
PBIP = REPO + '/raw/refs/heads/main/public/powerbi/CRM_ANA_PowerBI.pbip.zip'
PBI_GUIDE = REPO + '/blob/main/docs/POWER_BI_SIN_COSTO.md'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = 'CRM ANA · Guía interactiva de instalación y uso'
prs.core_properties.subject = 'Capacitación del equipo · CRM web 1.3.0 y extensión 1.2.1'
prs.core_properties.author = 'Equipo de Coordinación · CRM ANA'
prs.core_properties.keywords = 'ANA, CRM, instalación, tareas, agenda, capacitación, interactiva'
slides, meta, pending, link_log = {}, [], [], []

def rgb(color):
    return RGBColor.from_string(C.get(color, color))

def box(s, x, y, w, h, fill='white', line=None, radius=True):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                            Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line); sh.line.width = Pt(.8)
    else:
        sh.line.fill.background()
    if radius:
        sh.adjustments[0] = .08
    return sh

def txt(s, text, x, y, w, h, size=18, color='ink', bold=False, align=None,
        valign=MSO_ANCHOR.TOP, margin=0, font='Arial'):
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    if font == 'Consolas': font = 'Courier New'
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.font.name = font; p.font.size = Pt(size)
        p.font.bold = bold; p.font.color.rgb = rgb(color)
        p.space_after = Pt(5); p.line_spacing = 1.12
        if align is not None: p.alignment = align
    return sh

def bind(sh, target=None, url=None, tip=None):
    if target:
        pending.append((sh, target, tip or 'Ir a ' + target))
    elif url:
        sh.click_action.hyperlink.address = url
        for node in sh._element.xpath('.//a:hlinkClick'):
            node.set('tooltip', tip or 'Abrir enlace')
        link_log.append(dict(kind='web', destination=url))
    return sh

def button(s, text, x, y, w, target=None, url=None, primary=True, h=.48):
    bg = box(s,x,y,w,h,'blue' if primary else 'white',None if primary else 'line')
    label = txt(s,text,x+.1,y+.015,w-.2,h-.03,13,'white' if primary else 'blue',True,
                align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
    for sh in (bg,label): bind(sh,target,url,text)
    return bg

def pill(s, text, x, y, w, fill='light', color='blue', target=None, size=11):
    bg=box(s,x,y,w,.31,fill)
    label=txt(s,text,x+.06,y+.015,w-.12,.28,size,color,True,
              align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
    if target:
        for sh in (bg,label): bind(sh,target,tip=text)

def pic(s, path, x, y, w, h, frame=False):
    path=Path(path)
    with Image.open(path) as im:
        iw,ih=im.size
    scale=min(w/iw,h/ih); dw,dh=iw*scale,ih*scale
    if frame: box(s,x-.04,y-.04,w+.08,h+.08,'white','line')
    sh=s.shapes.add_picture(str(path),Inches(x+(w-dw)/2),Inches(y+(h-dh)/2),Inches(dw),Inches(dh))
    for nv in sh._element.xpath('./p:nvPicPr/p:cNvPr'):
        nv.set('descr','Logotipo ANA original' if path==LOGO else 'Captura real del CRM con datos de demostración: '+path.stem)
    return sh

def note(s, text):
    s.notes_slide.notes_text_frame.text = text

def callout(s, title, body, x, y, w=3.4, h=1.15, number=None, fill='light', target=None):
    shapes=[box(s,x,y,w,h,fill)]
    inset=.2
    if number is not None:
        shapes.append(box(s,x+.18,y+.2,.38,.38,'blue'))
        shapes.append(txt(s,str(number),x+.2,y+.215,.34,.32,14,'white',True,PP_ALIGN.CENTER))
        inset=.73
    shapes.append(txt(s,title,x+inset,y+.16,w-inset-.18,.34,17,'ink',True))
    shapes.append(txt(s,body,x+inset,y+.56,w-inset-.18,h-.60,13,'muted'))
    if target:
        for sh in shapes: bind(sh,target,tip='Ver '+title)
    return shapes

def step(s, n, title, body, x, y, w=5.3):
    box(s,x,y+.03,.41,.41,'blue')
    txt(s,str(n),x+.04,y+.065,.33,.31,14,'white',True,PP_ALIGN.CENTER)
    txt(s,title,x+.65,y,w-.65,.39,19,'ink',True)
    txt(s,body,x+.65,y+.47,w-.65,.66,15,'muted')

def warning(s, text, y=6.2, x=.55, w=12.23, yellow=False):
    box(s,x,y,w,.45,'yellow' if yellow else 'light')
    txt(s,text,x+.15,y+.07,w-.3,.33,12.5,'ochre' if yellow else 'blue')

def make(key, title, subtitle, section='USO DIARIO', dark=False):
    s=prs.slides.add_slide(prs.slide_layouts[6]); slides[key]=s
    s.background.fill.solid();s.background.fill.fore_color.rgb=rgb('navy' if dark else 'white')
    meta.append(dict(key=key,title=title,section=section,dark=dark))
    if not dark:
        box(s,0,0,W,.075,'blue',radius=False)
        txt(s,section,.55,.3,8,.25,10.5,'blue',True)
        pic(s,LOGO,10.95,.22,1.82,.8)
        txt(s,title,.55,.86,11.7,.66,32,'ink',True)
        txt(s,subtitle,.57,1.57,12.0,.53,15,'muted')
    return s

def screenshot_slide(key,title,subtitle,filename,items,section='USO DIARIO',foot=None):
    s=make(key,title,subtitle,section)
    pic(s,CAP/filename,.58,2.25,8.05,4.18,frame=True)
    pill(s,'CAPTURA · DATOS DE DEMOSTRACIÓN',.76,6.55,4.4,size=9)
    for i,(a,b) in enumerate(items):callout(s,a,b,8.95,2.24+i*1.4,3.8,1.35,number=i+1)
    if foot: note(s,foot)
    return s

# 01 · Cover
s=make('portada','CRM & Agenda Cloud','',dark=True,section='INICIO')
box(s,.62,.5,3.1,1.42,'white');pic(s,LOGO,.87,.65,2.58,1.12)
txt(s,'COORDINACIÓN INSTITUCIONAL',.7,2.25,7,.35,12,'C8DAF5',True)
txt(s,'Cada coordinación,\nen su espacio.',.67,2.87,7.2,1.65,43,'white',True)
txt(s,'Guía interactiva para instalar el CRM en las PCs\ny organizar el trabajo diario del equipo.',.72,4.85,7.1,.95,20,'DCE7F8')
button(s,'Comenzar  →',.73,6.12,2.45,target='menu',h=.57)
for i,(title,body) in enumerate([('01  Instala','Chrome + acceso del equipo'),('02  Coordina','Contactos, tareas y agenda'),('03  Da seguimiento','Avisos, reportes y respaldo')]):
    y=2.28+i*1.22
    box(s,8.65,y,3.85,1.0,'244A81')
    txt(s,title,8.9,y+.14,3.35,.35,20,'white',True)
    txt(s,body,8.9,y+.58,3.35,.3,13,'DAE7FB')
txt(s,'WEB 1.3.0  ·  EXTENSIÓN 1.2.1  ·  28 OCT. 2026',8.67,6.25,4,.38,10,'DAE7FB',True)
note(s,'Presentación de capacitación del CRM de coordinación del equipo. Utiliza el logotipo original proporcionado por el usuario, sin redibujarlo. Abrir en PowerPoint, Presentación con diapositivas → Desde el principio. Los botones azules y las tarjetas del menú son enlaces. No requiere macros ni complementos. Las capturas se generaron en la aplicación real con datos sintéticos sobre una D1 local aislada; no contienen contactos de producción ni el token. Las funciones documentadas corresponden a web 1.3.0 y lector 1.2.1.')

# 02 · Interactive hub
s=make('menu','Elige tu recorrido','Haz clic en una tarjeta. Puedes volver a este menú desde cualquier diapositiva.','NAVEGACIÓN INTERACTIVA')
for x,y,n,a,b,dest in [(.58,2.32,'01','Instalar en una PC','Descarga, Chrome, acceso y operador.','modos'),(6.92,2.32,'02','Usar el CRM','Contactos, tareas, reuniones y avisos.','directorio'),(.58,4.2,'03','Reportes y continuidad','Excel, Power BI, respaldo y actualizaciones.','reportes'),(6.92,4.2,'04','Practicar y resolver','Casos guiados, diagnóstico y autoevaluación.','ayuda')]:
    bg=box(s,x,y,5.82,1.55,'light');num=txt(s,n,x+.25,y+.2,.65,.6,29,'blue',True)
    t=txt(s,a,x+1.03,y+.23,4.45,.41,22,'ink',True)
    b=txt(s,b,x+1.03,y+.82,4.4,.48,15,'muted')
    for sh in (bg,num,t,b):bind(sh,dest,tip='Abrir recorrido: '+a)
button(s,'Explorar las 15 funciones  →',.65,6.15,3.8,target='funciones',primary=False)
txt(s,'Activa “Presentación con diapositivas” para usar todos los enlaces.',5.05,6.22,7.7,.4,13,'muted')
note(s,'Ofrecer al grupo elegir el recorrido. Para formación completa, seguir las flechas Siguiente; para consulta rápida, elegir tarjeta. Los enlaces internos funcionan en modo presentación. En edición, PowerPoint puede requerir Ctrl+clic o clic derecho → Abrir vínculo. Los enlaces a descargas o al CRM abren el navegador y requieren Internet. Esta presentación guía el uso; no ejecuta acciones en el CRM.')

# 03 · 15 linked features
s=make('funciones','Las 15 funciones acordadas','Cada tarjeta abre la explicación correspondiente. Los avisos entre operadores se incluyen además como herramienta de coordinación.','MAPA DE FUNCIONES')
features=[('Operador por PC','operador'),('Responsables asignados','ficha'),('Viendo ahora: otra PC','presencia'),('Espacios en la nube','ficha'),('Directorio y buscador','directorio'),('CUT y documentos','ficha'),('Bandejas de gestión','directorio'),('Texto a tarea','tareas'),('Tablero de pendientes','tareas'),('Bitácora interna','bitacora'),('Checklist por contacto','tareas'),('Agenda, minutas y avisos','agenda'),('Formatos y firma /','formatos'),('Excel y respaldo','reportes'),('Proyecto Power BI','powerbi')]
for i,(label,dest) in enumerate(features):
    col,row=divmod(i,5);x=.58+col*4.23;y=2.34+row*.77
    bg=box(s,x,y,3.98,.62,'light')
    n=txt(s,f'{i+1:02}',x+.15,y+.15,.4,.3,13,'blue',True)
    t=txt(s,label,x+.72,y+.13,3.05,.35,14.5,'ink',True)
    for sh in (bg,n,t):bind(sh,dest,tip='Ver función '+str(i+1)+': '+label)
button(s,'Avisar a otro operador  →',.65,6.38,3.7,target='avisos')
note(s,'Mapa basado en docs/REVISION_15_FUNCIONES.md, complementado por docs/NOVEDADES_1.3.md. Los comportamientos se muestran dentro del CRM, manteniendo WhatsApp intacto. Power BI se entrega como proyecto PBIP, por decisión del usuario, no como PBIX compilado. La apertura nativa de ese proyecto requiere validación en Windows.')

# 04 · Modes
s=make('modos','Una misma agenda, dos formas de entrar','La información se comparte en la nube: puedes continuar desde otra PC o desde el navegador.','INSTALAR EN CADA PC')
for x,title,subtitle,items in [(.6,'En la oficina','Chrome + WhatsApp Web',['Abre el CRM en el panel derecho.','El chat activo orienta la ficha del contacto.','Ajusta el panel a unos 380–450 px.']),(6.95,'Fuera de la oficina','Navegador de PC o celular',['Abre el enlace del CRM directamente.','Usa Directorio, Tareas, Agenda y reportes.','No necesitas tener WhatsApp Web abierto.'])]:
    box(s,x,2.4,5.78,3.6,'light')
    txt(s,title,x+.28,2.67,5.22,.48,25,'blue',True)
    txt(s,subtitle,x+.28,3.3,5.22,.4,18,'ink',True)
    for j,line in enumerate(items):txt(s,'• '+line,x+.28,4.0+j*.53,5.22,.48,16,'ink')
button(s,'Abrir CRM web  ↗',7.26,6.18,2.85,url=APP)
button(s,'Preparar la PC  →',.88,6.18,2.85,target='requisitos')
note(s,'El modo panel utiliza la extensión local; la web se carga desde Cloudflare Pages. El modo web funciona en navegadores modernos y no requiere la extensión. Ambos necesitan Internet para leer y guardar; no hay una cola de escrituras fuera de línea. Cada perfil de Chrome conserva su operador. La agenda no depende de que el celular maestro tenga guardado el número.')

# 05 · Requirements
s=make('requisitos','Antes de instalar: tres cosas listas','Este procedimiento se repite una sola vez en cada PC del equipo.','INSTALAR EN CADA PC')
for i,(title,body) in enumerate([('Chrome actualizado','El panel lateral requiere Chrome 116 o superior. En una PC administrada, consulta a TI si la carga está restringida.'),('Token del equipo','Solicita el TEAM_TOKEN al administrador por un canal privado. Se introduce dentro del CRM; no se publica en la PPT.'),('Una carpeta permanente','Guarda la extensión en una carpeta fija. Mantén los archivos allí después de instalarla.')]):
    x=.6+i*4.24
    box(s,x,2.5,3.96,3.25,'light')
    txt(s,f'0{i+1}',x+.25,2.75,3.4,.8,38,'blue',True)
    txt(s,title,x+.25,3.58,3.4,.8,22,'ink',True)
    txt(s,body,x+.25,4.55,3.4,1.16,15,'muted')
warning(s,'Instalación descomprimida: no necesitas pagar la inscripción de Chrome Web Store.',6.2)
note(s,'Para operadores no se requiere Node, npm, GitHub ni acceso administrativo a Cloudflare. Esas herramientas pertenecen al mantenimiento del proyecto. Instalar desde el ZIP de la extensión, no desde el ZIP del repositorio completo. La opción elegida por el usuario es sin pagos adicionales y no publica la extensión en la tienda. No intentar eludir restricciones de TI.')

# 06 · Download
s=make('descarga','1. Descarga y descomprime el paquete correcto','Usa crm-ana-extension.zip. El ZIP completo del proyecto contiene más carpetas y no se carga directamente.','INSTALAR EN CADA PC')
step(s,1,'Descarga el ZIP','Pulsa el botón de esta diapositiva y guárdalo en tu equipo.',.65,2.45,5.5)
step(s,2,'Descomprime todo','Windows: Extraer todo. Mac: doble clic sobre el archivo ZIP.',.65,3.68,5.5)
step(s,3,'Conserva la carpeta','Comprueba que manifest.json esté dentro de la carpeta que vas a seleccionar.',.65,4.91,5.5)
box(s,6.75,2.48,5.9,3.38,'navy')
txt(s,'CARPETA DE LA EXTENSIÓN',7.02,2.75,5.2,.35,11,'C4D8F4',True)
txt(s,'CRM-ANA-Extension\n  manifest.json\n  background.js\n  content_script.js\n  index.html\n  shell.js · shell.css · protocol.js\n  ana-logo.png · icons/',7.03,3.35,5.22,2.36,14.5,'white',font='Courier New')
button(s,'Descargar extensión 1.2.1  ↗',7.0,6.12,4.9,url=ZIP,h=.56)
note(s,'El enlace descarga artifacts/crm-ana-extension.zip desde la rama main del repositorio. El nombre CRM-ANA-Extension es una sugerencia para la carpeta local: puede usarse cualquier nombre y ubicación permanente. Dentro deben estar manifest.json y sus archivos hermanos, no otra carpeta intermedia. No seleccionar el ZIP. Si GitHub pide iniciar sesión, usar la cuenta con acceso al repositorio.')

# 07 · Chrome setup diagram
s=make('chrome','2. Carga la carpeta en Chrome','Escribe chrome://extensions en la barra de direcciones y presiona Enter.','INSTALAR EN CADA PC')
box(s,.65,2.4,8.2,3.85,'light','line')
box(s,.65,2.4,8.2,.55,'navy',radius=False)
txt(s,'chrome://extensions',.89,2.53,7.6,.3,15,'white',font='Consolas')
txt(s,'Extensiones',.95,3.25,4,.45,25,'ink',True)
txt(s,'Modo desarrollador',5.56,3.37,2.0,.32,13,'ink')
box(s,7.68,3.35,.63,.29,'blue');box(s,8.02,3.38,.21,.22,'white')
button(s,'Cargar descomprimida',1.0,4.07,3.34,target='abrir',h=.53)
box(s,1.0,4.93,7.42,.97,'white','line')
txt(s,'CRM & Agenda Cloud · Coordinación ANA',1.22,5.11,6.95,.33,18,'ink',True)
txt(s,'Versión 1.2.1  ·  Activada',1.22,5.56,6.9,.25,12,'blue')
callout(s,'Activa el interruptor','Modo desarrollador debe estar encendido.',9.18,2.55,3.5,1.15,number=1)
callout(s,'Elige la carpeta','Selecciona la que contiene manifest.json.',9.18,3.92,3.5,1.15,number=2)
callout(s,'Comprueba la tarjeta','La extensión debe aparecer activa, versión 1.2.1.',9.18,5.29,3.5,1.15,number=3)
txt(s,'Esquema orientativo de Chrome; los textos pueden variar según la versión.',.85,6.47,8.0,.28,10.5,'muted')
note(s,'La ilustración es un esquema editable, no una captura de una instalación real. En Chrome → Extensiones → Administrar extensiones, activar Modo desarrollador y pulsar Cargar descomprimida. Elegir la carpeta descomprimida. La tarjeta debe mostrar versión 1.2.1. Para una actualización ya instalada, conservar la carpeta anterior y sustituir archivos, luego Recargar. No eliminar el directorio mientras Chrome use la copia descomprimida.')

# 08 · Login
s=make('abrir','3. Abre el panel e ingresa al CRM','Abre web.whatsapp.com, inicia tu sesión de WhatsApp y pulsa el icono de la extensión.','INSTALAR EN CADA PC')
pic(s,CAP/'01-acceso.png',.9,2.15,3.12,4.68,frame=True)
step(s,1,'Fija y abre la extensión','En el menú de extensiones de Chrome, busca CRM ANA y fija su icono.',4.6,2.37,7.65)
step(s,2,'Introduce el token del equipo','El administrador te lo entrega por un canal privado. La dirección del CRM ya viene configurada.',4.6,3.75,7.65)
step(s,3,'Entra al espacio de trabajo','Pulsa el botón de acceso. Después selecciona quién está operando esta PC.',4.6,5.13,7.65)
note(s,'El contenedor 1.2.1 carga automáticamente el origen fijo https://crm-whatsapp-ana.pages.dev/panel. La pantalla actual solo pide el token del equipo; las capturas antiguas que pedían una URL corresponden a otra versión. El selector de operador aparece después de la autenticación. El token no se muestra ni se incluye en la presentación. Si el panel informa de almacenamiento limitado, se puede trabajar en esa sesión o usar Abrir para ir a una pestaña, donde se inicia sesión de nuevo.')

# 09 · Operator
s=make('operador','4. Elige quién opera esta PC','Las notas, tareas, reuniones y acuerdos registran al operador activo. Revisa tu nombre al comenzar.','PRIMER ACCESO')
names=['Hanz','Fernando','Acxel','Andrés','Shirley','Jaruska','Marco','Alejandra','Marcelo','Luciano','Sheila','Cristina']
for i,n in enumerate(names):
    col,row=i%3,i//3
    pill(s,n,.65+col*1.85,2.72+row*.69,1.65,fill='blue' if n=='Acxel' else 'light',color='white' if n=='Acxel' else 'blue',size=14)
txt(s,'OPERADORES INICIALES',.67,2.26,5.3,.32,11,'muted',True)
step(s,1,'Selecciona tu nombre','Usa “¿Quién opera esta PC?” en la barra superior.',6.65,2.5,5.8)
step(s,2,'Agrega nuevos integrantes','Pulsa +, escribe el nombre y elige Crear operador.',6.65,3.81,5.8)
step(s,3,'Gestiona cambios del equipo','Antes de eliminar a alguien, reasigna sus espacios y tareas pendientes.',6.65,5.12,5.8)
warning(s,'El token es compartido; el selector atribuye el trabajo al nombre elegido y no crea permisos individuales.',6.53)
note(s,'La selección se conserva por perfil de navegador; en una PC compartida cada persona debe cambiarla al empezar. Los operadores se administran en el botón + junto al selector y la lista se comparte con todas las PCs. Los nombres iniciales son los aprobados por el usuario. La autoría histórica permanece si se elimina un integrante. El selector es declarativo: cualquier poseedor del token tiene acceso al equipo, incluida la restauración; no hay cuentas personales ni permisos por rol.')

# 10 · Directory
screenshot_slide('directorio','Encuentra la coordinación en segundos','Directorio es el punto de partida para abrir una ficha, ver pendientes y priorizar el trabajo.','02-directorio.png',[
 ('Busca por contexto','Nombre, teléfono, CUT, órgano, texto de una nota o tarea.'),
 ('Filtra el directorio','Combina órgano/unidad y responsable asignado.'),
 ('Prioriza con bandejas','Urgentes, Pendientes, En elaboración, Para V°B°, Cerrado y Mis Coordinaciones.')],foot='Usar el Buscador global sobre Directorio. Las bandejas filtran estados de gestión de espacios; no cambian WhatsApp. Mis Coordinaciones utiliza el operador activo. También puede filtrarse por órgano y operador. El nombre Cerrado se presenta según el texto actual del control. Los KPIs de la pantalla resumen la agenda compartida.')

# 11 · Passive matching
s=make('chat','Abre un chat: el CRM muestra su contexto','El lector observa el nombre o teléfono del encabezado. El trabajo de coordinación ocurre dentro del panel.','USO CON WHATSAPP WEB')
box(s,.65,2.37,12.05,.77,'navy')
txt(s,'CHAT DE WHATSAPP  →  [Demo] Coordinación de Prensa',.95,2.59,11.45,.38,21,'white',True)
for i,(a,b,dest) in enumerate([('Ya existe una ficha','Si la coincidencia es única, se abre. También puedes usar Abrir ficha de este chat.','ficha'),('Es un contacto nuevo','Pulsa Crear espacio para este chat y completa su teléfono o ID conocido.','ficha'),('Hay nombres repetidos','Selecciona manualmente la ficha correcta desde Directorio.','directorio')]):
    x=.65+i*4.12
    box(s,x,3.58,3.82,2.4,'light')
    txt(s,a,x+.23,3.82,3.35,.78,21,'blue',True)
    txt(s,b,x+.23,4.73,3.35,1.06,15,'ink')
    button(s,'Ver cómo  →',x+.23,6.18,3.35,target=dest,primary=False)
note(s,'El lector es pasivo: no modifica el DOM de WhatsApp, no lee el historial de mensajes y no envía mensajes. Un encabezado con nombre no siempre proporciona el número real; el título de grupo no es un ID estable. Para evitar fichas equivocadas solo hay apertura automática ante coincidencia única. Registrar el número completo con código de país o el ID conocido; no inventar identificadores a partir de nombres. Si no cambia el contexto, seguir el diagnóstico de detección de chat.')

# 12 · Contact fields
s=make('ficha','Crea una ficha útil para todo el equipo','En Directorio pulsa Nuevo espacio o, desde el chat, Crear espacio para este chat.','ESPACIOS DE CONTACTO')
for x,title,rows in [(.65,'Identidad del contacto',['Nombre y cargo','Teléfono internacional o ID del grupo','Órgano / Unidad, oficina y entidad','Correo, anexo y ubicación']), (6.85,'Seguimiento de la gestión',['Estado y responsables del equipo','Código CUT y documento de referencia','Sede Central y Prensa ya están disponibles','Añadir órgano/unidad → Guardar unidad'])]:
    box(s,x,2.43,5.81,2.97,'light')
    txt(s,title,x+.24,2.67,5.3,.45,23,'blue',True)
    for j,r in enumerate(rows):txt(s,'• '+r,x+.25,3.36+j*.43,5.25,.4,16,'ink')
txt(s,'Guarda la ficha y elige una herramienta:',.72,5.76,11.9,.37,18,'ink',True)
for i,(label,dest) in enumerate([('Bitácora','bitacora'),('Tareas','tareas'),('Reuniones','reuniones'),('Formatos /','formatos')]):button(s,label,.72+i*3.15,6.3,2.88,target=dest)
note(s,'Campos obligatorios: nombre y teléfono o ID. Usar información real autorizada. Los contactos se guardan en D1 y no dependen de la libreta del celular maestro de WhatsApp. Los responsables pueden ser varios. Para una unidad nueva pulsar Añadir órgano/unidad, escribirla y Guardar unidad antes de Guardar espacio. Las opciones creadas se comparten con las otras PCs y quedan en el respaldo. Los CUT y documentos se asocian a la ficha y aparecen relacionados en los reportes.')

# 13 · Notes
s=make('bitacora','Bitácora: deja la coordinación por escrito','Registra hechos, próximos pasos y contexto que otro integrante necesitará retomar.','HERRAMIENTAS DEL CONTACTO')
pic(s,CAP/'03-bitacora.png',.7,2.2,4.18,4.55,frame=True)
step(s,1,'Abre Bitácora','Dentro de la ficha del contacto, escribe una actualización concreta.',5.4,2.38,6.92)
step(s,2,'Pulsa Registrar nota','El CRM agrega automáticamente autor, fecha y hora de Perú.',5.4,3.7,6.92)
step(s,3,'Elimina solo si corresponde','Pulsa Eliminar nota y confirma. Se borra ese registro para todo el equipo.',5.4,5.02,6.92)
note(s,'Ejemplo: “Se revisó el borrador. Jaruska enviará observaciones mañana y Acxel consolidará la respuesta”. Las notas no se editan desde la interfaz, pero pueden eliminarse con confirmación. Eliminar nota afecta solo ese registro; no elimina tareas, reuniones ni el resto del historial. La nota no es un mensaje enviado a WhatsApp. La captura usa datos de demostración y contiene un aviso interno registrado automáticamente.')

# 14 · Tasks
screenshot_slide('tareas','Una sección de Tareas, dos vistas','Lista y Tablero Kanban muestran los mismos registros. Cada tarea conserva el contacto y su CUT.','07-tareas.png',[
 ('Crea una tarea','Nueva tarea → contacto, título, responsable y vencimiento. Pega el texto de origen si lo necesitas.'),
 ('Actualiza su estado','Por hacer → En proceso → En revisión → Completado. Usa el selector de cada tarjeta.'),
 ('Revisa tu carga','Filtra por integrante o usa Mis tareas. En la ficha del contacto también está su checklist.')],foot='El tablero no depende de arrastrar tarjetas: el control implementado para avanzar es el selector Estado de tarea. Lista y Kanban comparten datos y filtros. Para un mensaje de WhatsApp, copiar el texto manualmente y pegarlo en el formulario del CRM. Para audios o documentos, escribir transcripción o referencia; el lector no los extrae automáticamente. La casilla del checklist marca Completado. Sin hora explícita, una tarea vence a las 17:00 de Perú.')

# 15 · Agenda
screenshot_slide('agenda','Agenda compartida: crea un evento nuevo','Usa Agenda → Nuevo evento. Puedes vincularlo a un contacto o dejarlo como evento general.','08-agenda.png',[
 ('Define cuándo','Título, inicio y fin en hora de Perú. El fin debe ser posterior al inicio.'),
 ('Completa el contexto','Descripción, participantes, lugar y enlace de videollamada opcional.'),
 ('Guarda y gestiona','Crear evento. Después puedes editarlo, cancelarlo, eliminarlo o exportarlo.')],foot='El calendario muestra reuniones de fichas y eventos generales. Elegir mes/día y buscar por evento, contacto o CUT. Los eventos nuevos conservan inicio y fin explícitos. Las reuniones previas siguen disponibles. Las acciones de edición y eliminación se sincronizan dentro del CRM. Después de exportar a un calendario personal, esos cambios se reflejan por separado en ese calendario.')

# 16 · Export calendar
s=make('calendarios','Llévalo a Google, Apple u Outlook','Abre el evento o la reunión guardada y elige cómo agregarlo a tu calendario personal.','AGENDA Y CALENDARIOS')
cards=[('Google Calendar','Pulsa Google Calendar.\nSe abre un borrador en tu cuenta.\nComprueba la hora y pulsa Guardar.'),('Apple / Outlook (.ics)','Descarga el archivo .ics.\nÁbrelo o impórtalo en tu calendario.\nConfirma el calendario de destino.'),('Varios eventos','Filtra el mes, el día o la búsqueda.\nPulsa Exportar selección .ics.\nImporta la selección visible.')]
for i,(a,b) in enumerate(cards):
    x=.65+i*4.13
    box(s,x,2.48,3.83,3.35,'light')
    txt(s,str(i+1),x+.24,2.74,3.35,.58,33,'blue',True)
    txt(s,a,x+.24,3.55,3.35,.75,22,'ink',True)
    txt(s,b,x+.24,4.39,3.35,1.2,15,'muted')
warning(s,'Es una copia al calendario: si cambias o eliminas el evento en el CRM, actualiza también tu calendario personal.',6.2)
note(s,'La exportación no vincula cuentas, no es sincronización bidireccional y no envía invitaciones. Google abre un borrador que la persona confirma con Guardar; los participantes se incluyen como texto descriptivo. ICS sirve para Apple, Outlook, Google y otros programas compatibles. El calendario receptor convierte la hora a su zona local; el CRM usa America/Lima. Reuniones antiguas sin hora final usan una hora de duración. Al importar repetidamente puede ser necesario gestionar duplicados. Los archivos ICS no incluyen TEAM_TOKEN.')

# 17 · Meetings
s=make('reuniones','De la reunión al acuerdo y a la tarea','Dentro del contacto abre Reuniones: agenda el encuentro y registra lo que debe ocurrir después.','HERRAMIENTAS DEL CONTACTO')
for i,(a,b) in enumerate([('Agendar reunión','Tema, fecha/hora, participantes, sala o Meet/Zoom y estado.'),('Añadir acuerdos','Escribe cada compromiso y define responsable y fecha.'),('Dar seguimiento','Crear tareas de acuerdos agrega los compromisos al tablero.')]):
    x=.65+i*4.13
    box(s,x,2.44,3.83,2.28,'light')
    txt(s,f'0{i+1}',x+.23,2.68,3.3,.51,28,'blue',True)
    txt(s,a,x+.23,3.35,3.35,.43,21,'ink',True)
    txt(s,b,x+.23,4.0,3.35,.61,15,'muted')
box(s,.65,5.05,12.06,1.29,'navy')
txt(s,'Generar minuta',.92,5.27,3.1,.4,23,'white',True)
txt(s,'Prepara el resumen de la reunión en Formatos /.\nCopia el texto o abre un borrador en WhatsApp para revisarlo y enviarlo tú.',4.45,5.28,7.87,.8,17,'white')
note(s,'Registrar tanto los acuerdos de texto como los acuerdos estructurados con responsable y fecha si se necesitan tareas. Crear tareas de acuerdos evita duplicados para la misma reunión. Si los acuerdos ya tienen tareas vinculadas, gestionar el avance desde Tareas. Generar minuta lleva el texto a Formatos /; Copiar o Abrir borrador en WhatsApp registran la preparación en la bitácora. La persona confirma cualquier envío de WhatsApp. Los eventos generales se crean en Agenda; las reuniones con acuerdos se gestionan en la ficha.')

# 18 · Templates
s=make('formatos','Formatos /: responde con orden y firma','Los textos y enlaces se comparten en la nube; cada operador puede preparar su propia respuesta.','HERRAMIENTAS DEL CONTACTO')
for i,(a,b) in enumerate([('/drive','Carpetas y enlaces'),('/tdr','Textos reutilizables'),('/cut','Consultas de seguimiento')]):
    x=.65+i*2.05
    box(s,x,2.52,1.83,1.19,'light')
    txt(s,a,x+.17,2.73,1.5,.4,21,'blue',True)
    txt(s,b,x+.17,3.22,1.5,.4,11,'muted')
step(s,1,'Selecciona un comando','Carga el formato y ajusta el contenido al contacto.',.68,4.06,6.05)
step(s,2,'Activa la firma si corresponde','Ejemplo: *Acxel - Coord. ANA*.',.68,5.36,6.05)
pic(s,CAP/'06-formatos.png',7.2,2.24,5.15,4.5,frame=True)
note(s,'Los formatos iniciales son editables y no sustituyen los enlaces reales del equipo. El comando /drive inicial apunta al servicio general; el administrador debe registrar la carpeta real cuando corresponda. Se puede buscar comando, seleccionar formato, modificar texto y activar la firma automática. Crear o editar formato compartido permite nuevos comandos. Copiar y Abrir borrador en WhatsApp no envían mensajes automáticamente; la bitácora registra texto copiado o borrador abierto, no entrega confirmada.')

# 19 · Internal notices
s=make('avisos','“Jaruska, este contacto quiere hablar contigo”','Envía el aviso dentro del CRM y conserva el vínculo a la ficha del contacto.','AVISOS ENTRE OPERADORES')
for x,title in [(.65,'1 · Acxel envía desde la ficha'),(6.84,'2 · Jaruska confirma desde la campana')]:
    box(s,x,2.36,5.83,3.93,'light')
    txt(s,title,x+.22,2.57,5.39,.4,18,'blue',True)
pic(s,CAP/'04-enviar-aviso.png',.88,3.14,5.36,2.92)
pic(s,CAP/'11-recibir-aviso.png',7.08,3.14,5.36,2.92)
txt(s,'Avisar a otro operador → destinatario → mensaje → Enviar aviso interno.',.76,6.48,5.6,.48,12,'muted')
txt(s,'Recibidos → Abrir espacio → Recibido. El remitente lo ve en Enviados.',6.95,6.48,5.6,.48,12,'muted')
note(s,'Desde la ficha de un contacto pulsar Avisar a otro operador. El nombre del remitente es el operador activo; elegir destinatario y escribir qué necesita el contacto. El aviso se guarda y deja una nota en la bitácora. La campana del destinatario muestra Recibidos, Enviados y Recordatorios. Marcar leído y Recibido son estados distintos: Recibido confirma atención al remitente. Si Jaruska no tiene abierto el CRM, el aviso queda guardado y lo verá al entrar. Mientras el CRM está abierto la bandeja consulta novedades cada cinco segundos. No es un mensaje de WhatsApp, correo o SMS.')

# 20 · Presence and reminders
s=make('presencia','Coordina sin duplicar esfuerzos','El CRM muestra quién está viendo la misma ficha y recuerda lo próximo a vencer.','EQUIPO Y RECORDATORIOS')
box(s,.65,2.48,6.0,3.63,'light')
txt(s,'Viendo ahora',.92,2.74,5.4,.48,26,'blue',True)
box(s,.92,3.52,5.46,1.1,'yellow')
txt(s,'Jaruska también tiene abierto\neste espacio en otra PC',1.1,3.73,5.08,.87,20,'ochre',True)
txt(s,'Coordinen quién registra cada cambio.\nLa cinta avisa; no bloquea la edición.',.97,5.02,5.3,.8,17,'ink')
for y,big,title,body in [(2.55,'15 min','Antes de una reunión','Revisa la campana y abre el contacto.'),(3.81,'1 hora','Antes de vencer una tarea','También se muestran las vencidas.'),(5.07,'17:00','Si la tarea no tiene hora','Se usa el horario de Perú.')]:
    box(s,7.0,y,5.68,1.02,'light')
    txt(s,big,7.22,y+.2,1.5,.52,25,'blue',True)
    txt(s,title,8.92,y+.16,3.48,.34,17,'ink',True)
    txt(s,body,8.92,y+.58,3.48,.34,12.5,'muted')
warning(s,'Los recordatorios aparecen mientras el CRM está abierto y hay un operador seleccionado.',6.42)
note(s,'La presencia envía un latido cada cuatro segundos y expira aproximadamente a los quince segundos sin latido. Otra pestaña del mismo perfil se distingue por dispositivo; dos perfiles pueden aparecer como dos PCs. La alerta es informativa, no bloqueo. Los recordatorios de tareas y reuniones se calculan mientras la aplicación está abierta; se pueden descartar para ese operador en esa PC. Los avisos internos durables son distintos de los recordatorios en pantalla.')

# 21 · Reports
s=make('reportes','Centro de datos: del seguimiento al reporte','Elige período e integrante. Descarga el Excel maestro o un PDF para revisión y reuniones.','REPORTES Y CONTINUIDAD')
box(s,.65,2.44,6.16,3.89,'light')
txt(s,'5 hojas en un solo Excel',.93,2.73,5.6,.48,25,'blue',True)
for i,row in enumerate(['1_Resumen_KPIs','2_Agenda_Espacios','3_Tareas_y_CUTs','4_Bitacora_Historica','5_Reuniones_Acuerdos']):
    box(s,.95,3.5+i*.47,5.57,.36,'white')
    txt(s,row,1.11,3.54+i*.47,5.2,.3,15,'ink',font='Consolas')
step(s,1,'Define qué vas a revisar','Semana, mes, rango personalizado o historial; todo el equipo o un integrante.',7.17,2.5,5.25)
step(s,2,'Descarga el reporte','Descargar .xlsx o Descargar reporte PDF.',7.17,3.88,5.25)
step(s,3,'Explora el libro','Incluye cinco tablas y tres tablas dinámicas nativas para reorganizar en Excel.',7.17,5.15,5.25)
note(s,'Los filtros del Centro de datos afectan Excel/PDF, no el respaldo JSON. Las tareas pendientes anteriores al período se conservan para no ocultar trabajo abierto. El libro contiene cinco tablas estructuradas y tres tablas dinámicas nativas con cachés: tareas por operador/estado, coordinaciones por órgano/estado y reuniones por contacto/estado. El formato de hoja y filtros pueden usarse en Excel. La generación y estructura OOXML se verificaron; la apertura/reorganización con Excel instalado se comprueba en los equipos del usuario.')

# 22 · Power BI
s=make('powerbi','Power BI: analiza el CRM desde Windows','Usa Power BI Desktop para trabajo local. El proyecto preparado se conecta a la API o a un Excel descargado.','REPORTES Y CONTINUIDAD')
box(s,.65,2.48,5.8,3.54,'navy')
txt(s,'Uso local sin Power BI Pro',.95,2.79,5.18,.56,25,'white',True)
txt(s,'1  Instala Power BI Desktop en Windows.\n2  Extrae todo el ZIP del proyecto.\n3  Abre CRM_ANA.pbip.\n4  Configura los parámetros y Actualizar.',.97,3.69,5.12,1.8,18,'white')
box(s,6.86,2.48,5.8,3.54,'light')
txt(s,'Parámetros de conexión',7.14,2.79,5.23,.46,23,'blue',True)
txt(s,'pOrigen: API\npBaseUrl: '+APP+'\npTeamToken: tu token privado del equipo',7.15,3.66,5.15,1.55,17,'ink')
txt(s,'El token se introduce en Power BI; mantenlo fuera de archivos que compartas públicamente.',7.15,5.32,5.15,.55,13,'muted')
button(s,'Descargar proyecto .pbip  ↗',.93,6.35,4.93,url=PBIP)
button(s,'Abrir guía paso a paso  ↗',7.15,6.35,4.93,url=PBI_GUIDE,primary=False)
note(s,'Power BI Desktop requiere Windows y se puede usar localmente sin contratar Pro. No hay versión nativa para Mac: usar una PC Windows existente del equipo; no se propone una VM de pago. El ZIP contiene Report, SemanticModel y el archivo CRM_ANA.pbip; conservar juntos. En Transformar datos → Administrar parámetros: pOrigen API, pBaseUrl https://crm-whatsapp-ana.pages.dev, pTeamToken valor privado del equipo, pArchivoExcel vacío. Si Power Query pide credenciales, usar Anónimo, pues la consulta M envía el token por cabecera X-Team-Token. Cerrar y aplicar y Actualizar. Para Excel elegir pOrigen Excel y ruta pArchivoExcel. Cinco páginas: Resumen ejecutivo, Tareas y CUT, Directorio, Reuniones y Bitácora. El proyecto se verifica estructuralmente y aún requiere abrirse y validarse en Power BI Desktop de Windows. Importación: refresco manual, no streaming. No publicar datos institucionales con Publicar en la web. No subir el proyecto ya configurado con token a GitHub.')

# 23 · Backup
s=make('respaldo','Entrega de cargo: conserva la continuidad','El respaldo JSON reúne los datos de negocio para recuperar o migrar el trabajo del equipo.','REPORTES Y CONTINUIDAD')
step(s,1,'Descarga una copia completa','Centro de datos → Descargar respaldo .json. Guárdala en una ubicación privada del equipo.',.7,2.45,6.05)
step(s,2,'Verifica antes de restaurar','Identifica qué copia usarás y coordina una pausa con los operadores.',.7,3.84,6.05)
step(s,3,'Restaura cuando corresponda','Restaurar desde .json → selecciona el archivo → escribe RESTAURAR → confirma.',.7,5.23,6.05)
box(s,7.1,2.5,5.5,3.85,'light')
txt(s,'La copia conserva',7.39,2.79,4.9,.47,25,'blue',True)
txt(s,'Contactos y CUT · notas · tareas\nReuniones, acuerdos y eventos\nFormatos · operadores · unidades\nAvisos internos y vínculos de tareas',7.42,3.6,4.82,1.5,17,'ink')
txt(s,'Restaurar reemplaza los datos actuales.\nLa interfaz descarga una copia previa.',7.42,5.45,4.82,.62,15,'ochre',True)
note(s,'El JSON v3 incluye los datos de negocio, no la cuenta Cloudflare, TEAM_TOKEN, dominio, UUID de D1 ni sesiones/presencia temporal. El nuevo administrador configura su instancia y restaura el archivo si se migra a otra cuenta. La restauración valida y reemplaza de forma atómica; la UI descarga un respaldo antes de ejecutarla. Solo usarla para recuperación o migración; no es una forma de actualizar el software. Límite de la interfaz 1,4 MB; para más volumen el administrador usa respaldo nativo de D1. Cualquier poseedor del token y operador seleccionado tiene acceso a restaurar; proteger la copia porque contiene información del equipo.')

# 24 · Updates
s=make('actualizar','Actualiza el CRM sin reinstalar cada cambio','La aplicación web y el lector de WhatsApp tienen versiones separadas.','ACTUALIZACIONES')
for x,title,tag,lines in [(.65,'Funciones del CRM','WEB 1.3.0',['Se publican desde GitHub a Cloudflare.','Si aparece Actualizar CRM, guarda tu trabajo.','Confirma la recarga para ver la nueva versión.']), (6.87,'Lector de WhatsApp','EXTENSIÓN 1.2.1',['Solo se reemplaza si cambia el lector o permisos.','Cierra el panel y sustituye los archivos de su carpeta.','Pulsa ↻ Recargar en chrome://extensions.'])]:
    box(s,x,2.47,5.8,3.73,'light')
    pill(s,tag,x+.25,2.74,2.25)
    txt(s,title,x+.25,3.32,5.25,.55,25,'ink',True)
    for j,line in enumerate(lines):txt(s,'• '+line,x+.25,4.12+j*.58,5.22,.61,15.5,'ink')
warning(s,'Después de recargar la extensión, recarga también WhatsApp Web con Ctrl+R (Windows) o ⌘R (Mac).',6.46)
note(s,'La web revisa nuevas versiones cada minuto mientras está visible y al volver a ella; no descarta borradores mediante recarga forzada. Guardar formularios antes de actualizar. Cambios del script del encabezado, permisos, puente o dominio requieren nuevo paquete. Como el usuario eligió no pagar la inscripción, Chrome Web Store no se ha publicado y las actualizaciones del lector se distribuyen manualmente. Mantener la carpeta anterior conserva su identidad descomprimida. Nunca borrar D1 ni restaurar JSON para actualizar código. Un token no cambia por instalar una versión; solo el administrador puede rotarlo en Cloudflare.')

# 25 · Help hub
s=make('ayuda','Si algo no funciona, elige el síntoma','Usa el recorrido adecuado antes de reinstalar o modificar datos.','AYUDA INTERACTIVA')
for i,(title,body,dest) in enumerate([('No carga la extensión','Falta el manifiesto o Chrome dice que no puede leerlo.','ayuda-manifest'),('El chat no cambia','Seleccionas otro contacto y el panel mantiene el contexto anterior.','ayuda-chat'),('No puedo entrar / API','Token inválido, error de conexión o respuesta que no es JSON.','ayuda-api')]):
    x=.65+i*4.12
    callout(s,title,body,x,2.65,3.83,2.4,number=i+1,target=dest)
    button(s,'Resolver  →',x+.2,5.45,3.43,target=dest)
button(s,'Ir al ejercicio práctico  →',4.62,6.36,4.1,target='practica',primary=False)
note(s,'Los botones llevan a soluciones concretas para los problemas que aparecieron durante la puesta en marcha. No se requiere borrar datos ni crear otra base para estos diagnósticos. Si el equipo tiene restricciones administrativas de Chrome, consultar a TI. El formulario correcto de despliegue es Pages; el operador no debe usar wrangler deploy ni configurar un Worker para abrir el CRM.')

# 26 · Manifest problem
s=make('ayuda-manifest','“Falta el archivo de manifiesto”','Chrome está buscando manifest.json en la carpeta que seleccionaste.','RESOLVER: INSTALACIÓN')
step(s,1,'Descomprime el ZIP de la extensión','No selecciones el archivo .zip ni la carpeta del repositorio completo.',.75,2.5,7.0)
step(s,2,'Abre la carpeta y mira su contenido','Debes ver manifest.json junto a background.js e index.html.',.75,3.91,7.0)
step(s,3,'Selecciona esa carpeta exacta','Vuelve a Cargar descomprimida y confirma que aparezca la tarjeta del CRM.',.75,5.32,7.0)
box(s,8.28,2.6,4.1,2.77,'navy')
txt(s,'CARPETA CORRECTA',8.56,2.9,3.54,.3,12,'DAE7FB',True)
txt(s,'CRM-ANA-Extension\n  manifest.json\n  background.js\n  index.html',8.55,3.59,3.54,1.65,18,'white',font='Consolas')
button(s,'Volver al paso de Chrome  →',8.3,5.75,4.05,target='chrome')
note(s,'En el ZIP exclusivo de la extensión manifest.json está en la raíz. La descarga del repositorio completo tiene una carpeta exterior y el código fuente, por lo que no debe cargarse su raíz como extensión. En Finder o Explorador puede haber una doble carpeta por el modo de extracción: abrir hasta llegar a manifest.json y seleccionar esa ubicación. Mantener el resto de archivos; no copiar únicamente el manifiesto.')

# 27 · Chat detection
s=make('ayuda-chat','El chat no se actualiza: revisa el lector','Primero comprueba que la tarjeta de la extensión muestre versión 1.2.1.','RESOLVER: CONTEXTO DEL CHAT')
for i,(a,b) in enumerate([('Recarga la extensión','Cierra el panel y pulsa ↻ Recargar en chrome://extensions.'),('Recarga WhatsApp Web','Ctrl+R en Windows o ⌘R en Mac. Después abre otra vez el panel.'),('Abre otro chat y vuelve a detectar','Usa Volver a detectar chat y comprueba el nombre mostrado.'),('Selecciona o crea la ficha','Sin coincidencia única, usa Directorio o Crear espacio para este chat.')]):
    x=.7+(i%2)*6.24;y=2.58+(i//2)*1.85
    callout(s,a,b,x,y,5.8,1.56,number=i+1)
warning(s,'El nombre guardado en WhatsApp puede no revelar el teléfono; verifica los datos antes de vincular una ficha.',6.37)
note(s,'Recargar solo el panel no reactiva el content script en una pestaña de WhatsApp que ya estaba abierta. Es indispensable recargar WhatsApp después de actualizar la extensión. El lector depende del encabezado disponible en WhatsApp y puede necesitar mantenimiento si Meta lo cambia. La búsqueda y apertura manual de Directorio siguen disponibles. Si hay dos fichas con el mismo nombre, el CRM evita abrir una incorrecta automáticamente.')

# 28 · API/access
s=make('ayuda-api','Acceso o API: identifica el mensaje','El operador revisa el acceso; el administrador revisa el despliegue cuando corresponde.','RESOLVER: CONEXIÓN')
for i,(a,b) in enumerate([('Token inválido / 401','Introduce nuevamente el token vigente. Si fue cambiado, solicita el nuevo al administrador.'),('No respondió JSON','Abre el CRM web y avisa al administrador para comprobar Pages Functions y la URL del despliegue.'),('Servicio no disponible / 503','El administrador debe revisar TEAM_TOKEN y el binding DB en Cloudflare. No borres la base.')]):
    callout(s,a,b,.72,2.45+i*1.28,11.87,1.1,number=i+1)
button(s,'Abrir CRM para comprobar  ↗',.94,6.48,4.4,url=APP)
txt(s,'Si hay error de red, recupera la conexión antes de volver a guardar.',5.77,6.53,6.7,.35,12.5,'muted')
note(s,'El código 401 corresponde a token inválido. Si falta un TEAM_TOKEN válido en el entorno, la API devuelve 503; consultar la configuración de producción y volver a desplegar después de cambiar secretos/bindings. Una respuesta HTML en vez de JSON puede significar funciones ausentes, URL/ruta incorrecta u otro fallo; no se atribuye una causa sin revisar. No subir solo dist con arrastrar y soltar, porque no instala Pages Functions. El panel actual usa un origen fijo y no requiere pegar URL. No compartir capturas del token ni de la URL privada de Power BI. Las escrituras fallidas no se consideran guardadas; verificar el resultado antes de repetir para evitar registros duplicados.')

# 29 · Practice
s=make('practica','Práctica guiada: una coordinación completa','Ejemplo: un contacto solicita conversar con Jaruska y entregar un material para revisión.','APRENDER HACIENDO')
practice=[('1','Abre o crea su ficha','ficha'),('2','Asigna CUT y responsables','ficha'),('3','Registra el contexto en Bitácora','bitacora'),('4','Envía un aviso a Jaruska','avisos'),('5','Crea la tarea con fecha límite','tareas'),('6','Agenda la reunión y exporta el evento','calendarios')]
for i,(n,t,dest) in enumerate(practice):
    col,row=i%2,i//2;x=.7+col*6.23;y=2.38+row*1.12
    bg=box(s,x,y,5.78,.86,'light');num=txt(s,n,x+.2,y+.2,.4,.42,22,'blue',True)
    tt=txt(s,t,x+.88,y+.23,4.6,.42,17,'ink',True)
    for sh in (bg,num,tt):bind(sh,dest,tip='Repasar: '+t)
button(s,'Comprobar lo aprendido  →',4.52,6.25,4.35,target='pregunta',h=.56)
note(s,'Ejercicio de capacitación. Si se practica en producción, acordarlo con el administrador y usar un contacto de prueba claramente identificado para no mezclar actividades reales. La presentación no crea registros por pulsar estas tarjetas: solo abre las instrucciones. Durante el ejercicio, Jaruska debe tener elegido su nombre para ver el aviso y confirmar Recibido. El tutor puede pedir que el grupo cambie entre Lista y Kanban para comprobar que es la misma tarea.')

# 30 · Branching question
s=make('pregunta','¿Cómo avisas a Jaruska desde el contacto?','Haz clic en la respuesta que usarías. La presentación te llevará a la explicación.','AUTOEVALUACIÓN INTERACTIVA')
for i,(letter,a,dest) in enumerate([('A','Compartir el token con el contacto para que busque a Jaruska.','reintentar'),('B','Abrir su ficha → Avisar a otro operador → Jaruska → Enviar aviso interno.','correcto'),('C','Crear otro operador llamado Jaruska para registrar el pedido.','reintentar')]):
    y=2.5+i*1.25
    bg=box(s,.8,y,11.72,.97,'light')
    badge=box(s,1.02,y+.22,.5,.5,'blue')
    let=txt(s,letter,1.07,y+.29,.4,.33,17,'white',True,PP_ALIGN.CENTER)
    ans=txt(s,a,1.86,y+.25,10.13,.62,20,'ink',True)
    for sh in (bg,badge,let,ans):bind(sh,dest,tip='Elegir respuesta '+letter)
note(s,'Respuesta correcta B. La A comparte una credencial del equipo con una persona que puede no estar autorizada; la C duplica una identidad existente y no crea un aviso para la bandeja de Jaruska. El ejercicio se resuelve mediante enlaces internos, sin macros, almacenamiento de puntuaciones ni envío de respuestas.')

# 31 · Correct
s=make('correcto','Respuesta correcta','',section='AUTOEVALUACIÓN',dark=True)
pill(s,'BIEN RESUELTO',.8,.8,2.6,fill='white',color='blue')
txt(s,'El aviso llega al CRM\nde la persona indicada.',.8,1.74,11.6,1.6,40,'white',True)
txt(s,'Jaruska abre la campana, revisa el contacto y confirma Recibido.\nTú puedes ver la confirmación en Enviados.',.85,3.65,11.3,1.1,24,'DFEAFB')
txt(s,'El pedido queda guardado aunque ella no esté conectada al momento de enviarlo.',.85,5.15,11.3,.65,18,'DFEAFB')
button(s,'Abrir recursos y empezar  →',.86,6.18,4.3,target='recursos',h=.57)
button(s,'Repasar avisos',5.58,6.18,2.7,target='avisos',primary=False,h=.57)
note(s,'Confirmar que el grupo distingue Leído de Recibido. La confirmación indica una acción dentro del CRM, no una llamada realizada ni un mensaje externo entregado. Cerrar con el enlace del CRM y el paquete de instalación.')

# 32 · Retry
s=make('reintentar','Revisa cómo llega el aviso','El destinatario ya existe como operador; el pedido debe llegar a su bandeja del CRM.','AUTOEVALUACIÓN INTERACTIVA')
box(s,.82,2.5,11.66,2.62,'light')
txt(s,'Parte desde la ficha del contacto.',1.18,2.86,10.92,.65,30,'blue',True)
txt(s,'Busca la acción que permite elegir un operador destinatario,\nescribir el pedido y dejar constancia en la coordinación.',1.2,3.92,10.65,.9,23,'ink')
button(s,'Volver a responder  →',1.18,5.77,4.45,target='pregunta',h=.58)
button(s,'Repasar cómo enviar un aviso',6.62,5.77,5.42,target='avisos',primary=False,h=.58)
note(s,'No se trata de dar acceso al contacto ni de crear otro miembro del equipo. Repasar Avisar a otro operador en la ficha. El selector de Jaruska determina qué bandeja ve esa persona. Volver a la pregunta para elegir B.')

# 33 · Resources
s=make('recursos','Listos para coordinar','Abre el recurso que necesitas. El CRM es el espacio de trabajo; esta presentación es la guía.','RECURSOS CLICABLES')
resources=[('Abrir CRM','Directorio, tareas, agenda y reportes.',APP),('Descargar extensión','Paquete descomprimido para Chrome · 1.2.1.',ZIP),('Descargar Power BI','Proyecto PBIP para Power BI Desktop en Windows.',PBIP),('Guía de Power BI sin costo','Parámetros, credenciales y actualización local.',PBI_GUIDE)]
for i,(a,b,url) in enumerate(resources):
    col,row=i%2,i//2;x=.72+col*6.24;y=2.43+row*1.6
    bg=box(s,x,y,5.78,1.32,'light');tt=txt(s,a+'  ↗',x+.23,y+.2,5.3,.4,22,'blue',True)
    bb=txt(s,b,x+.23,y+.8,5.3,.38,14,'muted')
    for sh in (bg,tt,bb):bind(sh,url=url,tip='Abrir '+a)
txt(s,'Rutina diaria: elige tu operador → revisa avisos → actualiza tareas → consulta la agenda.',.86,6.1,11.7,.48,18,'ink',True)
button(s,'Volver al menú',5.0,6.67,3.33,target='menu',primary=False,h=.4)
note(s,'Recursos del proyecto en GitHub y Cloudflare. Los enlaces de GitHub pueden requerir iniciar sesión si el repositorio cambia de visibilidad. Nunca solicitar pago de registro de Chrome Web Store para este procedimiento. El CRM y sus datos requieren Internet y token autorizado. Para mantener el material, ejecutar node scripts/presentation-captures.mjs y python3 scripts/create-training-presentation.py; ambos trabajan con demostración local para capturas y no consultan la D1 de producción.')

# Global navigation, notes metadata, and a manual (click-only) slideshow.
sections={'INSTALAR':'modos','USAR':'directorio','REPORTES':'reportes','AYUDA':'ayuda'}
custom_next={'pregunta':'recursos','correcto':'recursos','reintentar':'pregunta'}
custom_prev={'correcto':'pregunta','reintentar':'pregunta','recursos':'menu'}
for i,m in enumerate(meta):
    s=slides[m['key']];dark=m['dark']
    if m['key'] not in ('portada','correcto'):
        box(s,.56,7.11,12.2,.012,'line',radius=False)
        txt(s,'CRM ANA · GUÍA INTERACTIVA',.6,7.22,3.6,.17,8.2,'muted')
        for j,(label,dest) in enumerate(sections.items()):
            bind(txt(s,label,4.1+j*1.02,7.21,.93,.2,8.4,'blue',True,PP_ALIGN.CENTER),dest,tip='Ir a '+label.lower())
    # Controls have fixed positions across slides.
    color='white' if dark else 'blue'
    for label,x,w,dest in [('‹',9.05,.35,custom_prev.get(m['key'],meta[max(0,i-1)]['key'])),('Menú',9.63,.78,'menu'),('›',10.7,.35,custom_next.get(m['key'],meta[min(len(meta)-1,i+1)]['key']))]:
        sh=txt(s,label,x,7.17,w,.24,12,color,True,PP_ALIGN.CENTER)
        bind(sh,dest,tip='Anterior' if label=='‹' else 'Siguiente' if label=='›' else 'Volver al menú')
    txt(s,f'{i+1:02} / {len(meta)}',11.55,7.21,1.2,.2,9,'DAE7FB' if dark else 'muted',align=PP_ALIGN.RIGHT)
    transition=OxmlElement('p:transition');transition.set('advClick','1');s._element.append(transition)
    # Explicit flat effects avoid inheriting Office's default gray drop shadows.
    for sh in s.shapes:
        sp_pr=getattr(sh._element,'spPr',None)
        if sp_pr is not None:
            sp_pr.append(OxmlElement('a:effectLst'))
        for ref in sh._element.xpath('./p:style/a:effectRef'):
            ref.set('idx','0')
    note(s,s.notes_slide.notes_text_frame.text+'\n\nFuentes: README.md; docs/NOVEDADES_1.3.md; docs/REVISION_15_FUNCIONES.md; docs/POWER_BI_SIN_COSTO.md; componentes React de la versión 1.3.0. Navegación: Menú y flechas en el pie; tarjetas y botones clicables. Guía elaborada el 28 de octubre de 2026 (Perú).')

for sh,target,tip in pending:
    sh.click_action.target_slide=slides[target]
    for node in sh._element.xpath('.//a:hlinkClick'):node.set('tooltip',tip)
    link_log.append(dict(kind='slide',destination=target))

output=OUT/'CRM_ANA_Guia_Interactiva.pptx'
prs.save(output)
manifest={'title':prs.core_properties.title,'file':output.name,'slides':len(meta),
          'web_version':'1.3.0','extension_version':'1.2.1','date_peru':'2026-10-28',
          'slides_index':[dict(number=i+1,**m) for i,m in enumerate(meta)],
          'internal_links':sum(x['kind']=='slide' for x in link_log),
          'external_links':sum(x['kind']=='web' for x in link_log),
          'logo_sha256':hashlib.sha256(LOGO.read_bytes()).hexdigest(),
          'pptx_sha256':hashlib.sha256(output.read_bytes()).hexdigest(),
          'screenshots':'Datos sintéticos en una D1 local aislada; sin secretos ni registros de producción.'}
(ROOT/'docs/capacitacion/contenido.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
(OUT/(output.name+'.sha256')).write_text(manifest['pptx_sha256']+'  '+output.name+'\n')
print(json.dumps({k:manifest[k] for k in ['file','slides','internal_links','external_links','pptx_sha256']},ensure_ascii=False))
