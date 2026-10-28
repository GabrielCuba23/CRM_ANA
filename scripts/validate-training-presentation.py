#!/usr/bin/env python3
"""Validate the downloadable presentation and its PDF reading copy."""
from pathlib import Path
import hashlib
import json
import posixpath
import re
import zipfile
from urllib.parse import urlsplit
from lxml import etree
from pptx import Presentation
import fitz

ROOT=Path(__file__).resolve().parents[1]
PPTX=ROOT/'artifacts/CRM_ANA_Guia_Interactiva.pptx'
PDF=ROOT/'artifacts/CRM_ANA_Guia_Interactiva.pdf'
manifest=json.loads((ROOT/'docs/capacitacion/contenido.json').read_text())
presentation=Presentation(PPTX)
assert len(presentation.slides)==manifest['slides']==33
index={x['key']:x['number'] for x in manifest['slides_index']}
graph={i:set() for i in range(1,34)}
external=set()
clicks=0
NS={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main'}

with zipfile.ZipFile(PPTX) as z:
    assert z.testzip() is None
    names=set(z.namelist())
    assert not any(re.search(r'vba|activex|embeddings/',name,re.I) for name in names)
    assert any(z.read(name)==(ROOT/'src/assets/ana-logo.png').read_bytes() for name in names if name.startswith('ppt/media/'))
    for name in names:
        if name.endswith(('.xml','.rels')): etree.fromstring(z.read(name))
    for i,slide in enumerate(presentation.slides,1):
        assert slide.notes_slide.notes_text_frame.text.strip(),f'Faltan notas: {i}'
        for shape in slide.shapes:
            assert shape.left>=0 and shape.top>=0,(i,shape.name)
            assert shape.left+shape.width<=presentation.slide_width+2,(i,shape.name)
            assert shape.top+shape.height<=presentation.slide_height+2,(i,shape.name)
        rels={r.get('Id'):r for r in etree.fromstring(z.read(f'ppt/slides/_rels/slide{i}.xml.rels'))}
        xml=etree.fromstring(z.read(f'ppt/slides/slide{i}.xml'))
        for link in xml.xpath('//a:hlinkClick',namespaces=NS):
            clicks+=1
            rel=rels[link.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')]
            target=rel.get('Target')
            if rel.get('Type').endswith('/slide'):
                destination=posixpath.normpath(posixpath.join('ppt/slides',target))
                assert destination in names,(i,destination)
                graph[i].add(int(re.search(r'slide(\d+)\.xml$',destination)[1]))
            else:
                assert rel.get('Type').endswith('/hyperlink')
                url=urlsplit(target)
                assert url.scheme=='https' and url.hostname in {'github.com','crm-whatsapp-ana.pages.dev'}
                assert 'token=' not in target.lower()
                external.add(target)
        if i==index['pregunta']:
            for prefix,dest in [('Compartir el token','reintentar'),('Abrir su ficha','correcto'),('Crear otro operador','reintentar')]:
                answer=next(sh for sh in slide.shapes if sh.has_text_frame and sh.text.startswith(prefix))
                assert answer.click_action.target_slide==presentation.slides[index[dest]-1]
    reached=set();queue=[index['menu']]
    while queue:
        current=queue.pop()
        if current in reached:continue
        reached.add(current);queue.extend(graph[current]-reached)
    assert len(reached)==33,('Diapositivas sin ruta desde el menú',set(graph)-reached)

pdf=fitz.open(PDF)
assert len(pdf)==33
assert all(len(page.get_text().strip())>150 for page in pdf)
for path in [PPTX,PDF]:
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    (path.parent/(path.name+'.sha256')).write_text(digest+'  '+path.name+'\n')
report={'slides':33,'all_slides_reachable':True,'clickable_shapes':clicks,'unique_external_urls':len(external),
        'quiz_branches_verified':True,'original_logo_embedded':True,'shape_bounds':'OK','xml_and_zip':'OK',
        'speaker_notes':33,'pdf_pages':len(pdf),'macros':False,
        'files':{p.name:{'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in [PPTX,PDF]},
        'native_powerpoint_validation':'No ejecutada en este entorno; render verificado mediante LibreOffice.'}
(ROOT/'docs/capacitacion/verificacion.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(report,ensure_ascii=False))
